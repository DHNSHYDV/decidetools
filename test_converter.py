import os
import shutil
import tempfile
import unittest
from pathlib import Path
import docx
from converter import docx_to_pdf, pdf_to_docx, convert_file
from app import app


class TestConverter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_dir = tempfile.mkdtemp(prefix="converter_test_")
        cls.sample_docx = os.path.join(cls.test_dir, "sample.docx")

        # Create a sample docx
        doc = docx.Document()
        doc.add_heading("Automated Test Document", 0)
        doc.add_paragraph("Hello world! This is a test for local file conversion.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"
        doc.save(cls.sample_docx)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.test_dir, ignore_errors=True)

    def test_1_docx_to_pdf(self):
        pdf_out = docx_to_pdf(self.sample_docx)
        self.assertTrue(os.path.isfile(pdf_out))
        self.assertTrue(pdf_out.endswith(".pdf"))
        self.assertGreater(os.path.getsize(pdf_out), 0)

    def test_2_pdf_to_docx(self):
        pdf_in = str(Path(self.sample_docx).with_suffix(".pdf"))
        docx_out = os.path.join(self.test_dir, "converted_back.docx")
        pdf_to_docx(pdf_in, docx_out)
        self.assertTrue(os.path.isfile(docx_out))
        self.assertGreater(os.path.getsize(docx_out), 0)

        # Verify that docx can be parsed by python-docx
        doc = docx.Document(docx_out)
        text = "\n".join([p.text for p in doc.paragraphs])
        self.assertIn("Automated Test Document", text)

    def test_3_convert_file_dispatcher(self):
        pdf_path = convert_file(self.sample_docx)
        self.assertTrue(os.path.exists(pdf_path))
        roundtrip_docx = convert_file(pdf_path, os.path.join(self.test_dir, "roundtrip.docx"))
        self.assertTrue(os.path.exists(roundtrip_docx))

    def test_4_flask_health_endpoint(self):
        client = app.test_client()
        res = client.get("/api/health")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json.get("status"), "ok")

    def test_5_flask_convert_docx_to_pdf(self):
        client = app.test_client()
        with open(self.sample_docx, "rb") as f:
            res = client.post("/api/convert", data={"file": (f, "test_doc.docx")})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, "application/pdf")
        self.assertGreater(len(res.data), 0)

    def test_6_flask_convert_pdf_to_docx(self):
        client = app.test_client()
        pdf_in = str(Path(self.sample_docx).with_suffix(".pdf"))
        with open(pdf_in, "rb") as f:
            res = client.post("/api/convert", data={"file": (f, "test_doc.pdf")})
        self.assertEqual(res.status_code, 200)
        self.assertIn("wordprocessingml", res.mimetype)
        self.assertGreater(len(res.data), 0)


if __name__ == "__main__":
    unittest.main()
