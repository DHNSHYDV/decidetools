import os
import shutil
import tempfile
import unittest
from pathlib import Path
import docx

from app import app, SITE_URL
from modules.file_converter import docx_to_pdf, pdf_to_docx
from modules.qr_studio import generate_qr_base64, generate_qr_image, build_wifi_payload
from modules.media_converter import format_duration, sanitize_filename


class TestDecidePlatform(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_dir = tempfile.mkdtemp(prefix="platform_test_")
        cls.sample_docx = os.path.join(cls.test_dir, "sample.docx")

        doc = docx.Document()
        doc.add_heading("Decide Platform Document", 0)
        doc.add_paragraph("Testing SaaS unified document conversion workflow.")
        doc.save(cls.sample_docx)

        cls.client = app.test_client()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.test_dir, ignore_errors=True)

    # 1. Canonical SEO Page Routes
    def test_canonical_pages_render_successfully(self):
        canonical_pages = [
            "/",
            "/tools/youtube-downloader",
            "/tools/youtube-to-mp3",
            "/tools/youtube-to-mp4",
            "/tools/qr-code-generator",
            "/tools/docx-to-pdf",
            "/tools/pdf-to-docx",
            "/about",
            "/privacy",
            "/terms",
        ]
        for page in canonical_pages:
            res = self.client.get(page)
            self.assertEqual(res.status_code, 200, f"Page {page} failed to load (status {res.status_code})")
            html = res.data.decode("utf-8")
            self.assertIn("Decide Solutions", html, f"Page {page} missing brand name")
            self.assertIn("<title>", html, f"Page {page} missing <title>")
            self.assertIn('name="description"', html, f"Page {page} missing meta description")
            self.assertIn('rel="canonical"', html, f"Page {page} missing canonical link")
            self.assertIn("<h1", html, f"Page {page} missing h1 heading")
            self.assertIn("application/ld+json", html, f"Page {page} missing JSON-LD schema")

    # 2. Legacy Route 301 Permanent Redirects
    def test_legacy_301_redirects(self):
        redirect_map = {
            "/tools/media": "/tools/youtube-downloader",
            "/tools/qr-code": "/tools/qr-code-generator",
            "/tools/converter": "/tools/docx-to-pdf",
        }
        for legacy_url, target_url in redirect_map.items():
            res = self.client.get(legacy_url, follow_redirects=False)
            self.assertEqual(res.status_code, 301, f"Legacy URL {legacy_url} did not return 301")
            self.assertTrue(res.headers.get("Location", "").endswith(target_url), 
                            f"Legacy URL {legacy_url} did not redirect to {target_url}")

    # 3. SEO Infrastructure: robots.txt and sitemap.xml
    def test_robots_txt(self):
        res = self.client.get("/robots.txt")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.content_type.startswith("text/plain"))
        content = res.data.decode("utf-8")
        self.assertIn("User-agent: *", content)
        self.assertIn("Allow: /", content)
        self.assertIn("Disallow: /api/", content)
        self.assertIn("Sitemap:", content)

    def test_sitemap_xml(self):
        res = self.client.get("/sitemap.xml")
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.content_type.startswith("application/xml"))
        content = res.data.decode("utf-8")
        self.assertIn("<urlset", content)
        self.assertIn("/tools/youtube-downloader</loc>", content)
        self.assertIn("/tools/youtube-to-mp3</loc>", content)
        self.assertIn("/tools/youtube-to-mp4</loc>", content)
        self.assertIn("/tools/qr-code-generator</loc>", content)
        self.assertIn("/tools/docx-to-pdf</loc>", content)
        self.assertIn("/tools/pdf-to-docx</loc>", content)
        self.assertIn("/about</loc>", content)
        self.assertIn("/privacy</loc>", content)
        self.assertIn("/terms</loc>", content)

    def test_google_site_verification(self):
        res = self.client.get("/google33d1629be034105c.html")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"google-site-verification: google33d1629be034105c.html", res.data)

    def test_custom_404(self):
        res = self.client.get("/this-route-does-not-exist-at-all")
        self.assertEqual(res.status_code, 404)
        html = res.data.decode("utf-8")
        self.assertIn("404 — Page Not Found", html)
        self.assertIn("Return to Home", html)

    def test_api_health(self):
        res = self.client.get("/api/health")
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertEqual(data["status"], "ok")
        self.assertEqual(data["platform"], "Decide Group Of Solutions")

    # 4. Document Studio
    def test_document_conversion(self):
        pdf_out = docx_to_pdf(self.sample_docx)
        self.assertTrue(os.path.isfile(pdf_out))

        docx_out = os.path.join(self.test_dir, "roundtrip.docx")
        pdf_to_docx(pdf_out, docx_out)
        self.assertTrue(os.path.isfile(docx_out))

    def test_api_convert_document(self):
        with open(self.sample_docx, "rb") as f:
            res = self.client.post("/api/convert", data={"file": (f, "test.docx")})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, "application/pdf")

    # 5. QR Code Studio
    def test_wifi_payload_builder(self):
        payload = build_wifi_payload("OfficeWiFi", "SecretPass123", "WPA")
        self.assertEqual(payload, "WIFI:T:WPA;S:OfficeWiFi;P:SecretPass123;H:false;;")

    def test_qr_generation_png_and_svg(self):
        png_bytes = generate_qr_image("https://decidegroup.com", output_format="png")
        self.assertGreater(len(png_bytes), 0)
        self.assertTrue(png_bytes.startswith(b"\x89PNG"))

        svg_bytes = generate_qr_image("https://decidegroup.com", output_format="svg")
        self.assertGreater(len(svg_bytes), 0)
        self.assertIn(b"<svg", svg_bytes)

    def test_api_qr_preview(self):
        res = self.client.post("/api/qr/preview", json={
            "data": "https://decidegroup.com",
            "fill_color": "#112233",
            "back_color": "#ffffff"
        })
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data["data_url"].startswith("data:image/png;base64,"))

    def test_api_qr_download(self):
        res = self.client.post("/api/qr/download", json={
            "data": "https://decidegroup.com",
            "format": "svg"
        })
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, "image/svg+xml")

    # 6. Media Converter Helpers
    def test_media_utilities(self):
        self.assertEqual(format_duration(65), "01:05")
        self.assertEqual(format_duration(3665), "01:01:05")
        self.assertEqual(sanitize_filename('Test/File:Name"<>|'), "TestFileName")

    def test_api_media_info_validation(self):
        res = self.client.post("/api/media/info", json={"url": ""})
        self.assertEqual(res.status_code, 400)


if __name__ == "__main__":
    unittest.main()
